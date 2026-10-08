import io
from decimal import Decimal

from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy
from django.views.decorators.http import require_POST

from core.models import AuditLog, client_ip
from ledger.services import PostingError
from users.permissions import require_perm

from . import services
from .forms import BandFormSet, NewRunForm, PayForm, RemitForm, SettingsForm
from .models import DelegationBand, PayrollRun, PayrollSettings, Payslip, Remittance

ZERO = Decimal("0")

# The deduction columns of the extract, in the order the NOSI and the Tax Authority expect them.
DEDUCTIONS = [
    ("si_employee", gettext_lazy("Social insurance — employee")),
    ("si_employer", gettext_lazy("Social insurance — employer")),
    ("emergency_fund", gettext_lazy("Emergency fund")), ("income_tax", gettext_lazy("Salary tax")),
    ("loan_deduction", gettext_lazy("Loan recovery")), ("other_deductions", gettext_lazy("Medical & other")),
    ("absence_deduction", gettext_lazy("Absence")),
]


def _journal(run):
    """The accrual entry lines, or [] when the chart is missing a payroll account (shown as a message)."""
    try:
        return services.journal_lines(run) if run.payslips.exists() else []
    except PostingError:
        return []


def _run(request, pk):
    return get_object_or_404(PayrollRun, pk=pk, company=request.company)


@require_perm("payroll.view")
def runs(request):
    today = timezone.localdate()
    form = NewRunForm(request.POST or None, initial={"year": today.year, "month": today.month})
    if request.method == "POST":
        if not request.membership.has_perm("payroll.create"):
            messages.error(request, _("You do not have permission to do that."))
        elif form.is_valid():
            year, month = form.cleaned_data["year"], form.cleaned_data["month"]
            run, created = PayrollRun.objects.get_or_create(company=request.company, year=year, month=month)
            if created:
                services.compute_run(run)
                AuditLog.record(request.company, request.user, "payroll.create", run, run.period_label,
                                ip=client_ip(request))
                messages.success(request, _("Payroll %(p)s calculated for %(n)s employees.") % {
                    "p": run.period_label, "n": run.payslips.count()})
            else:
                messages.info(request, _("A run for this month already exists."))
            return redirect("payroll:run", pk=run.pk)
    rows = [(r, r.totals()) for r in PayrollRun.objects.filter(company=request.company)[:36]]
    return render(request, "payroll/runs.html", {"form": form, "rows": rows})


@require_perm("payroll.view")
def run(request, pk):
    run = _run(request, pk)
    company = request.company
    totals = run.totals()
    s = PayrollSettings.for_company(company)
    today = timezone.localdate()
    remitted = {r.kind: r for r in run.remittances.select_related("journal_entry", "bank_account")}
    amounts = services.remittance_amounts(run)
    remit_rows = [(k, label, amounts[k], remitted.get(k),
                   RemitForm(company=company, initial={"kind": k, "pay_date": today, "bank_account": s.pay_bank_account}))
                  for k, label in Remittance.KINDS]
    return render(request, "payroll/run.html", {
        "run": run, "t": totals, "payslips": run.payslips.select_related("employee", "employee__department"),
        "signoffs": services.required_signoffs(company, totals["net_pay"]),
        "pay_form": PayForm(company=company, initial={"pay_date": today, "bank_account": s.pay_bank_account,
                                                      "cash_account": s.pay_cash_account}),
        "remit_rows": remit_rows, "deductions": DEDUCTIONS,
        "cash_total": sum((p.net_pay for p in run.payslips.select_related("employee") if p.employee.payment_method == "cash"), ZERO)})


ACTION_PERMS = {"recompute": "payroll.create", "submit": "payroll.create", "approve": "payroll.approve",
                "authorise": "payroll.authorise", "reject": "payroll.approve", "pay": "payroll.post",
                "remit": "payroll.post", "delete": "payroll.edit"}


@require_POST
@require_perm("payroll.view")
def run_action(request, pk):
    run = _run(request, pk)
    action, m = request.POST.get("action"), request.membership
    if action not in ACTION_PERMS or not m.has_perm(ACTION_PERMS[action]):
        messages.error(request, _("You do not have permission to do that."))
        return redirect("payroll:run", pk=pk)
    try:
        if action == "recompute":
            services.compute_run(run)
            messages.success(request, _("Recalculated."))
        elif action == "submit":
            services.submit(run, request.user)
            messages.success(request, _("Sent to the Finance Manager for approval."))
        elif action == "approve":
            services.approve(run, request.user, m)
            run.refresh_from_db()
            messages.success(request, _("Approved and posted to the ledger.") if run.status == "authorised"
                             else _("Approved. Waiting for the Managing Director."))
        elif action == "authorise":
            services.authorise(run, request.user, m)
            messages.success(request, _("Authorised. The payroll journal is posted."))
        elif action == "reject":
            services.reject(run, request.user, request.POST.get("reason", ""))
            messages.success(request, _("Returned to HR."))
        elif action == "pay":
            form = PayForm(request.POST, company=request.company)
            if not form.is_valid():
                messages.error(request, " ".join(e for errs in form.errors.values() for e in errs))
            else:
                d = form.cleaned_data
                services.mark_paid(run, request.user, d["pay_date"], d["bank_account"], d["cash_account"])
                messages.success(request, _("Salaries marked paid and posted."))
        elif action == "remit":
            form = RemitForm(request.POST, company=request.company)
            if not form.is_valid():
                messages.error(request, " ".join(e for errs in form.errors.values() for e in errs))
            else:
                d = form.cleaned_data
                services.remit(run, d["kind"], request.user, d["pay_date"], d["bank_account"], d["reference"])
                messages.success(request, _("Payment over posted."))
        elif action == "delete":
            services.delete_run(run, request.user)
            messages.success(request, _("Draft run deleted."))
            return redirect("payroll:runs")
    except services.PayrollError as exc:
        messages.error(request, str(exc))
    return redirect("payroll:run", pk=pk)


@require_perm("payroll.view")
def payslip(request, pk):
    p = get_object_or_404(Payslip.objects.select_related("run", "employee", "employee__department"),
                          pk=pk, run__company=request.company)
    return render(request, "payroll/payslip.html", {"p": p, "run": p.run})


@require_perm("payroll.view")
def payslips_print(request, pk):
    run = _run(request, pk)
    return render(request, "payroll/payslips_print.html", {
        "run": run, "slips": run.payslips.select_related("employee", "employee__department")})


@require_perm("payroll.view")
def deductions(request, pk):
    run = _run(request, pk)
    slips = list(run.payslips.select_related("employee", "employee__department"))
    if request.GET.get("format") == "xlsx":
        if not request.membership.has_perm("payroll.export"):
            messages.error(request, _("You do not have permission to do that."))
            return redirect("payroll:deductions", pk=pk)
        AuditLog.record(request.company, request.user, "payroll.export", run, run.period_label, ip=client_ip(request))
        return _excel(run, slips)
    return render(request, "payroll/deductions.html", {
        "run": run, "slips": slips, "deductions": DEDUCTIONS, "t": run.totals(),
        "lines": _journal(run)})


def _excel(run, slips):
    """One workbook: deductions, NOSI schedule, salary tax, payment schedule and the GL journal."""
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    def sheet(ws, title, header, rows, money_from=0):
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
        for col in ws.iter_cols(min_col=money_from + 1, min_row=2):
            for c in col:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00"
        for i, _h in enumerate(header, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18 if i > 2 else 26
        ws.freeze_panes = "C2"

    def e(p):
        return [p.employee.code, p.employee.name_en]

    sheet(wb.active, "Deductions", ["Code", "Employee", "Department", "Gross"] + [str(l) for _k, l in DEDUCTIONS] +
          ["Total deductions", "Net pay"],
          [e(p) + [p.employee.department.name_en if p.employee.department else "", p.gross] +
           [getattr(p, k) for k, _l in DEDUCTIONS] + [p.total_deductions, p.net_pay] for p in slips], 3)
    sheet(wb.create_sheet(), "NOSI", ["Code", "Employee", "National ID", "Insurance no.", "Insurable wage",
                                      "Employee 11%", "Employer 18.75%", "Emergency fund 1%", "Total"],
          [e(p) + [p.employee.national_id, p.employee.insurance_no, p.insurable_wage, p.si_employee, p.si_employer,
                   p.emergency_fund, p.si_employee + p.si_employer + p.emergency_fund]
           for p in slips if p.si_employee or p.si_employer], 4)
    sheet(wb.create_sheet(), "Salary tax", ["Code", "Employee", "National ID", "Gross", "Insurance (employee)",
                                            "Salary tax"],
          [e(p) + [p.employee.national_id, p.gross, p.si_employee, p.income_tax] for p in slips], 3)
    sheet(wb.create_sheet(), "Payments", ["Code", "Employee", "Paid by", "Bank", "Account / IBAN / InstaPay", "Net pay"],
          [e(p) + [p.employee.get_payment_method_display(), p.employee.bank_name, p.employee.bank_account, p.net_pay]
           for p in slips], 5)
    lines = _journal(run)
    sheet(wb.create_sheet(), "GL journal", ["Account", "Account name", "Cost centre", "Debit", "Credit", "Description"],
          [[l.account.code, l.account.name_en, str(l.cost_center or ""), l.debit, l.credit, str(l.description)]
           for l in lines if l.debit or l.credit], 3)
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="payroll-{run.period_label}-deductions.xlsx"'
    return response


@require_perm("payroll.edit")
def settings_view(request):
    company = request.company
    s = PayrollSettings.for_company(company)
    form = SettingsForm(request.POST or None, instance=s, company=company, prefix="s")
    bands = BandFormSet(request.POST or None, queryset=DelegationBand.objects.filter(company=company), prefix="bands")
    if request.method == "POST" and form.is_valid() and bands.is_valid():
        form.save()
        for obj in bands.save(commit=False):
            obj.company = company
            obj.save()
        for obj in bands.deleted_objects:
            obj.delete()
        AuditLog.record(company, request.user, "payroll.settings", s, "", ip=client_ip(request))
        messages.success(request, _("Saved."))
        return redirect("payroll:settings")
    return render(request, "payroll/settings.html", {"form": form, "bands": bands})


@require_perm("payroll.view")
def leave_provision(request):
    today = timezone.localdate()
    try:
        year, month = int(request.GET.get("year", today.year)), int(request.GET.get("month", today.month))
        assert 1 <= month <= 12
    except (ValueError, AssertionError):
        year, month = today.year, today.month
    rows, totals = services.leave_provision(request.company, year, month)
    return render(request, "payroll/leave_provision.html", {"rows": rows, "totals": totals, "year": year, "month": month,
                                                            "months": range(1, 13)})
