"""Batch C: reports page and new standard reports, Excel export, dashboard customising and drill-down,
the Settings gear pages, menu bookmarks and alerts."""
import io
import json
from datetime import timedelta

import pytest
from django.urls import reverse
from django.utils import timezone, translation
from openpyxl import load_workbook

from banking.models import BankTransfer, Cheque
from banking.services import create_bank_account, post_transfer
from contacts.models import Customer
from core.models import Bookmark, Feedback, UserPreference
from inventory import services as stock
from inventory.models import Item, Warehouse
from ledger.models import TaxRate
from ledger.services import Line, account_balance, post_journal, system_account
from reports import catalog, extra
from sales import services as sales
from sales.models import CustomerPayment, Invoice, InvoiceLine

from .conftest import D, acct


@pytest.fixture
def books(company, owner):
    """Opening cash, a bank, one overdue and one current invoice, a customer payment, a transfer and an asset bought."""
    today = timezone.localdate()
    bank = create_bank_account(company, kind="bank", currency="EGP", name_en="CIB", user=owner)
    cash = company.bank_accounts.get(kind="cash")
    post_journal(company, today - timedelta(days=60), [Line(bank.gl_account, debit=D(500000)),
                                                       Line(acct(company, "3100"), credit=D(500000))],
                 memo="Opening", source="opening", user=owner)
    wh = Warehouse.objects.get(company=company, code="MAIN")
    item = Item.objects.create(company=company, sku="SUP-1", name_en="Bur set", sales_price=D(1000), purchase_cost=D(600))
    stock.receive(item, wh, 20, D(600), today - timedelta(days=50), "OPEN")
    customer = Customer.objects.create(company=company, code="CUS-1", name_en="Nile Smile", currency_id="EGP")
    vat = TaxRate.objects.get(company=company, rate=14)
    invoices = []
    for days_ago, due_in in ((40, -10), (3, 27)):
        inv = Invoice.objects.create(company=company, customer=customer, date=today - timedelta(days=days_ago),
                                     due_date=today + timedelta(days=due_in), currency_id="EGP", warehouse=wh, created_by=owner)
        InvoiceLine.objects.create(document=inv, item=item, qty=D(2), unit_price=D(1000), tax_rate=vat)
        sales.post_invoice(inv, owner)
        invoices.append(inv)
    pay = CustomerPayment(company=company, date=today - timedelta(days=2), customer=customer, currency_id="EGP",
                          amount=D(2280), method="bank", bank_account=bank, created_by=owner)
    sales.post_payment(pay, [(invoices[1], D(2280))], owner)
    post_transfer(BankTransfer(company=company, date=today - timedelta(days=1), from_account=bank, to_account=cash,
                               amount=D(10000), amount_to=D(10000), created_by=owner), owner)
    post_journal(company, today - timedelta(days=1), [Line(system_account(company, "fixed_assets"), debit=D(50000)),
                                                      Line(bank.gl_account, credit=D(50000))],
                 memo="Dental chair", source="manual", user=owner)
    return {"bank": bank, "cash": cash, "customer": customer, "invoices": invoices, "today": today}


def test_cash_flow_statement_reconciles_to_cash(company, books):
    today = books["today"]
    start = today - timedelta(days=30)
    with translation.override("en"):
        rows, foot, info = extra.cash_flows(company, start, today)
    cash_now = account_balance(books["bank"].gl_account) + account_balance(books["cash"].gl_account)
    assert foot[1] == cash_now
    assert info["change"] == D(2280) - D(50000)  # payment in, chair out; the transfer is internal
    investing = [r for r in rows if r["cells"][0] == "Net cash from investing activities"]
    assert investing and investing[0]["cells"][1] == D(-50000)


def test_dashboard_cash_flow_ignores_transfers(company, books):
    from dashboard.services import cash_flow_block

    flow = cash_flow_block(company, books["today"])
    assert flow["today"] == account_balance(books["bank"].gl_account) + account_balance(books["cash"].gl_account)
    assert flow["money_in"] == D(500000) + D(2280)
    assert flow["money_out"] == D(50000)
    assert len(flow["chart"]["urls"]) == 12 and "/reports/cash-flows/" in flow["chart"]["urls"][-1]


@pytest.mark.parametrize("lang", ["en", "ar"])
def test_reports_page_and_every_report_open(client, company, owner, books, lang):
    client.force_login(owner)
    client.cookies["django_language"] = lang
    page = client.get(reverse("reports:index"))
    assert page.status_code == 200
    html = page.content.decode()
    assert 'id="repSearch"' in html and "data-fav-list" in html
    sections, _favs = catalog.visible_sections(company.memberships.get(user=owner))
    assert len(sections) >= 10
    for s in sections:
        for r in s["items"]:
            assert client.get(r["url"]).status_code == 200, r["key"]
    for extra_url in ("reports:cash_overview",):
        assert client.get(reverse(extra_url)).status_code == 200
    for period in ("this_month", "last_quarter", "last_year", "last_365"):
        assert client.get(reverse("reports:pl_compare"), {"period": period, "compare": "previous_year"}).status_code == 200


def test_customer_balance_reports_match_ageing(company, books):
    from sales.services import ar_aging

    as_of = books["today"]
    rows, foot = extra.balance_summary(extra.customer_open_items(company, as_of), "contacts:customer")
    assert foot[2] == sum(r["total"] for r in ar_aging(company, as_of))


def test_favourites(client, company, owner):
    client.force_login(owner)
    r = client.post(reverse("reports:favourite"), {"key": "pl_month"})
    assert r.json() == {"on": True}
    html = client.get(reverse("reports:index")).content.decode()
    fav_block = html.split("data-fav-list")[1].split("</details>")[0]
    assert reverse("reports:pl_month") in fav_block
    assert client.post(reverse("reports:favourite"), {"key": "pl_month"}).json() == {"on": False}
    assert client.post(reverse("reports:favourite"), {"key": "nope"}).status_code == 400


def test_export_to_excel(client, company, owner, books):
    client.force_login(owner)
    r = client.get(reverse("reports:export"), {"url": reverse("reports:sales_customer") + "?period=last_365"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    values = [c for row in wb.active.iter_rows(values_only=True) for c in row if c is not None]
    assert "Sales by customer" in values and "Nile Smile" in values
    assert 4000.0 in values  # two invoices of 2,000 net
    for bad in ("//evil.example/x", "https://evil.example/", "/reports/export/?url=/"):
        assert client.get(reverse("reports:export"), {"url": bad}).status_code == 404


def test_export_respects_the_target_permission(client, company, make_user):
    user = make_user("sales@scigate.test", "Sales officer")
    client.force_login(user)
    assert client.get(reverse("reports:export"), {"url": reverse("reports:payroll_summary")}).status_code == 403


def test_dashboard_layout_save_and_reset(client, company, owner, books):
    client.force_login(owner)
    layout = [{"key": "ar", "w": 2}, {"key": "cashflow", "w": 4}, {"key": "stock", "hidden": True}, {"key": "hack", "w": 9}]
    assert client.post(reverse("dashboard:layout"), {"layout": json.dumps(layout)}).json() == {"ok": True}
    pref = UserPreference.for_user(owner, company)
    assert [w["key"] for w in pref.dashboard_layout] == ["ar", "cashflow", "stock"]
    html = client.get(reverse("dashboard:home")).content.decode()
    assert html.index('data-widget="ar"') < html.index('data-widget="cashflow"') < html.index('data-widget="pl"')
    assert 'data-widget="ar" data-w="2"' in html and 'data-widget="cashflow" data-w="4"' in html
    assert 'data-widget="stock" data-w="1" aria-label="Stock alerts" hidden' in html
    client.post(reverse("dashboard:layout"), {"reset": "1"})
    pref.refresh_from_db()
    assert pref.dashboard_layout == []
    assert client.post(reverse("dashboard:layout"), {"layout": "not json"}).status_code == 400


def test_dashboard_drill_links_and_feed(client, company, owner, books):
    client.force_login(owner)
    html = client.get(reverse("dashboard:home")).content.decode()
    for target in (reverse("sales:invoices") + "?pay=overdue", reverse("banking:cheques") + "?due=7",
                   reverse("banking:cheques") + "?status=in_safe", reverse("reports:cash_overview"),
                   reverse("reports:ar_aging"), "/reports/cash-flows/?period=last_365"):
        assert target in html, target
    assert "Overdue invoices" in html and "Customer payments" in html  # business feed cards
    overdue = client.get(reverse("sales:invoices"), {"pay": "overdue"}).content.decode()
    assert books["invoices"][0].number in overdue and books["invoices"][1].number not in overdue
    assert "Clear filter" in overdue


def test_cheque_list_filters(client, company, owner):
    today = timezone.localdate()
    for n, due in (("1001", 3), ("1002", 30)):
        Cheque.objects.create(company=company, direction="in", number=n, amount=D(100), currency_id="EGP",
                              due_date=today + timedelta(days=due), issue_date=today, status="in_safe", party_name="X")
    client.force_login(owner)
    html = client.get(reverse("banking:cheques"), {"due": "7"}).content.decode()
    assert "#1001" in html and "#1002" not in html
    assert "#1002" in client.get(reverse("banking:cheques"), {"q": "1002"}).content.decode()


@pytest.mark.parametrize("lang", ["en", "ar"])
def test_settings_gear_pages(client, company, owner, lang):
    client.force_login(owner)
    client.cookies["django_language"] = lang
    home = client.get(reverse("dashboard:home")).content.decode()
    for name in ("core:settings", "users:users", "core:report_settings", "ledger:accounts", "core:additional_info",
                 "core:lists", "inventory:items", "core:recurring", "core:attachments", "banking:rules",
                 "banking:reconcile", "core:audit", "core:subscription", "core:feedback", "core:privacy", "core:cookies",
                 "core:accountant", "core:help", "core:alerts"):
        assert reverse(name) in home, name
        assert client.get(reverse(name)).status_code == 200, name
    assert client.get(reverse("core:alerts"), {"partial": "1"}).status_code == 200


def test_settings_actions(client, company, owner):
    client.force_login(owner)
    client.post(reverse("core:report_settings"), {"period": "this_month", "compare": "previous_year"})
    r = client.get(reverse("reports:profit_loss"))
    assert r.context["from"] == timezone.localdate().replace(day=1)
    client.post(reverse("core:additional_info"), {"industry": "medical", "website": "scigate.example", "founded": "2009"})
    company.refresh_from_db()
    assert company.profile["industry"] == "medical" and company.profile["website"].startswith("https://")
    client.post(reverse("core:feedback"), {"kind": "idea", "message": "Add budgets"})
    assert Feedback.objects.get().message == "Add budgets"
    client.post(reverse("core:privacy"), {"blur": "on"})
    assert 'data-priv-default="1"' in client.get(reverse("dashboard:home")).content.decode()
    accountant = Membership_role_id(company)
    assert f"?role={accountant}" in client.get(reverse("core:accountant")).content.decode()
    form = client.get(reverse("users:user_new"), {"role": accountant}).context["form"]
    assert [r.pk for r in form.initial["roles"]] == [accountant]


def Membership_role_id(company):
    from users.models import Role
    return Role.objects.get(company=company, name_en="Accountant").pk


def test_menu_bookmark_toggle(client, company, owner):
    client.force_login(owner)
    url = reverse("ledger:accounts")
    Bookmark.objects.filter(user=owner, url=url).delete()
    assert client.post(reverse("core:bookmark_toggle"), {"url": url, "title": "Chart"}).json() == {"on": True}
    html = client.get(reverse("dashboard:home")).content.decode()
    assert f'class="bm on" type="button" data-bm="{url}"' in html
    assert client.post(reverse("core:bookmark_toggle"), {"url": url}).json() == {"on": False}
    assert client.post(reverse("core:bookmark_toggle"), {"url": "https://evil.example/"}).status_code == 400
    accounting = client.get(reverse("dashboard:home")).content.decode()
    assert reverse("core:accountant") in accounting
