"""Batch G: approval limits for purchase orders and bills, own records only, warehouse restriction, costs and margins,
the Delete permission and the "Who can do what" report."""
import pytest
from django.urls import reverse
from django.utils import timezone

from contacts.models import Customer, Supplier
from crm.models import Lead
from inventory.models import Item, Warehouse
from inventory.services import receive
from ledger.models import Account, TaxRate
from purchases import services as purchasing
from purchases.models import Bill, BillLine, PurchaseOrder, PurchaseOrderLine
from sales.models import Invoice
from users.models import ApprovalLimit, Membership, Role

from .conftest import D


@pytest.fixture
def g(company, owner):
    today = timezone.localdate()
    supplier = Supplier.objects.create(company=company, code="SUP-1", name_en="Medent", currency_id="EGP")
    main = Warehouse.objects.get(company=company, code="MAIN")
    alx = Warehouse.objects.create(company=company, code="ALX", name_en="Alexandria", name_ar="الإسكندرية")
    item = Item.objects.create(company=company, sku="MCH-1", name_en="Dental chair", type="inventory",
                               sales_price=D(100000), purchase_cost=D(60000))
    expense = Account.objects.filter(company=company, type="expense", is_group=False, is_active=True).first()
    return {"today": today, "supplier": supplier, "main": main, "alx": alx, "item": item, "expense": expense,
            "vat": TaxRate.objects.get(company=company, rate=14)}


def _membership(user):
    return Membership.objects.get(user=user)


def _po(company, owner, g, amount):
    po = PurchaseOrder.objects.create(company=company, supplier=g["supplier"], date=g["today"], currency_id="EGP",
                                      warehouse=g["main"], created_by=owner)
    PurchaseOrderLine.objects.create(document=po, item=g["item"], qty=D(1), unit_price=D(amount))
    purchasing.submit_order(po, owner)
    return po


# ---------- 1. Approval limits ----------
def test_purchase_order_passed_up_above_limit(company, owner, make_user, g):
    fm = make_user("fm@scigate.test", "Finance manager")      # purchase order limit 1,000,000
    md = make_user("md@scigate.test", "Managing director")    # no limit
    po = _po(company, owner, g, 1_500_000)
    assert purchasing.approve_order(po, fm, _membership(fm)) == "passed_up"
    po.refresh_from_db()
    assert po.status == "submitted" and len(po.approvals) == 1
    with pytest.raises(purchasing.PurchaseError):            # cannot approve twice
        purchasing.approve_order(po, fm, _membership(fm))
    assert [o.pk for o in purchasing.orders_waiting_for(company, md, _membership(md))] == [po.pk]
    assert purchasing.approve_order(po, md, _membership(md)) == "approved"
    po.refresh_from_db()
    assert po.status == "open" and po.approved_by == md and len(po.approvals) == 2

    small = _po(company, owner, g, 200_000)
    assert purchasing.approve_order(small, fm, _membership(fm)) == "approved"


def test_approver_without_a_limit_is_refused(company, owner, make_user, g):
    role = Role.objects.create(company=company, name_en="Buyer lead", permissions=["purchases.view", "purchases.approve"])
    user = make_user("buyer@scigate.test", "Purchasing officer")
    _membership(user).roles.add(role)
    po = _po(company, owner, g, 1000)
    from django.utils import translation

    with translation.override("en"):
        assert "no approval limit" in purchasing.order_approval_block(po, user, _membership(user))


def test_bill_above_limit_waits_for_approval(client, company, owner, make_user, g):
    accountant = make_user("acc@scigate.test", "Accountant")  # bill limit 500,000
    fm = make_user("fm@scigate.test", "Finance manager")      # bill limit 1,000,000
    bill = Bill.objects.create(company=company, supplier=g["supplier"], supplier_invoice_no="B-1", date=g["today"],
                               due_date=g["today"], currency_id="EGP", created_by=accountant)
    BillLine.objects.create(document=bill, account=g["expense"], description="Fit-out", qty=D(1), unit_price=D(600000))
    client.force_login(accountant)
    client.post(reverse("purchases:bill_action", args=[bill.pk]), {"action": "post"})
    bill.refresh_from_db()
    assert bill.status == "draft" and bill.approval_requested_by == accountant
    assert purchasing.bills_waiting_for(company, accountant, _membership(accountant)) == []
    assert [b.pk for b in purchasing.bills_waiting_for(company, fm, _membership(fm))] == [bill.pk]
    client.force_login(fm)
    inbox = client.get(reverse("approvals:inbox")).content.decode()
    assert "B-1" in inbox
    client.post(reverse("purchases:bill_action", args=[bill.pk]), {"action": "post"})
    bill.refresh_from_db()
    assert bill.status == "posted" and bill.approval_requested_at is None


# ---------- 2. Own records only ----------
def test_own_records_only(client, company, owner, make_user, g):
    rep = make_user("rep@scigate.test", "Sales officer")
    other = make_user("rep2@scigate.test", "Sales officer")
    Role.objects.filter(company=company, name_en="Sales officer").update(own_records_only=True)
    mine = Customer.objects.create(company=company, code="C-1", name_en="Mine", currency_id="EGP", account_manager=rep)
    theirs = Customer.objects.create(company=company, code="C-2", name_en="Theirs", currency_id="EGP", account_manager=other)
    inv_mine = Invoice.objects.create(company=company, customer=mine, number="INV-M", date=g["today"], due_date=g["today"],
                                      currency_id="EGP", created_by=owner)
    inv_theirs = Invoice.objects.create(company=company, customer=theirs, number="INV-T", date=g["today"],
                                        due_date=g["today"], currency_id="EGP", created_by=owner)
    Lead.objects.create(company=company, name="My lead", owner=rep)
    Lead.objects.create(company=company, name="Their lead", owner=other)
    client.force_login(rep)
    page = client.get(reverse("sales:invoices"))
    assert [r["obj"].number for r in page.context["rows"]] == ["INV-M"]
    assert client.get(reverse("sales:invoice", args=[inv_theirs.pk])).status_code == 404
    assert client.get(reverse("sales:invoice", args=[inv_mine.pk])).status_code == 200
    customers = client.get(reverse("contacts:customers")).content.decode()
    assert "Mine" in customers and "Theirs" not in customers
    assert client.get(reverse("contacts:customer", args=[theirs.pk])).status_code == 404
    form = client.get(reverse("sales:invoice_new")).context["form"]
    assert list(form.fields["customer"].queryset) == [mine]
    leads = client.get(reverse("crm:list", args=["lead"])).content.decode()
    assert "My lead" in leads and "Their lead" not in leads
    client.post(reverse("contacts:customer_new"), {"name_en": "New one", "currency": "EGP", "payment_terms_days": 30,
                                                    "country": "Egypt", "credit_limit": 0, "is_active": "on"})
    assert Customer.objects.get(name_en="New one").account_manager == rep
    # A second, unrestricted role lifts the limit.
    _membership(rep).roles.add(Role.objects.get(company=company, name_en="Accountant"))
    client.force_login(rep)
    assert len(client.get(reverse("sales:invoices")).context["rows"]) == 2


# ---------- 3. Warehouses ----------
def test_warehouse_restriction(client, company, owner, make_user, g):
    keeper = make_user("store@scigate.test", "Storekeeper")
    _membership(keeper).warehouses.add(g["alx"])
    receive(g["item"], g["main"], D(2), D(60000), g["today"], "OPEN-1")
    receive(g["item"], g["alx"], D(3), D(60000), g["today"], "OPEN-2")
    client.force_login(keeper)
    rows = client.get(reverse("inventory:stock")).context["rows"]
    assert [(i.sku, q) for i, q, _v in rows] == [("MCH-1", D(3))]
    form = client.get(reverse("inventory:adjustment_new")).context["form"]
    assert list(form.fields["warehouse"].queryset) == [g["alx"]]
    transfer = client.get(reverse("inventory:transfer_new")).context["form"]
    assert list(transfer.fields["from_warehouse"].queryset) == [g["alx"]]
    assert set(transfer.fields["to_warehouse"].queryset) == {g["main"], g["alx"]}
    po_main = PurchaseOrder.objects.create(company=company, supplier=g["supplier"], date=g["today"], currency_id="EGP",
                                           warehouse=g["main"], created_by=owner)
    assert client.get(reverse("purchases:order", args=[po_main.pk])).status_code == 404


# ---------- 4. Costs and margins ----------
def test_costs_hidden_without_permission(client, company, owner, make_user, g):
    seller = make_user("seller@scigate.test", "Sales officer")
    client.force_login(seller)
    items = client.get(reverse("inventory:items")).content.decode()
    assert "Purchase cost" not in items and "60,000" not in items
    assert client.get(reverse("reports:stock_detail")).status_code == 403
    home = client.get(reverse("dashboard:home")).content.decode()
    assert "Profit &amp; loss" not in home and "Bank accounts" not in home
    keeper = make_user("store@scigate.test", "Storekeeper")
    client.force_login(keeper)
    assert "Purchase cost" in client.get(reverse("inventory:items")).content.decode()


# ---------- 5. Who can do what ----------
def test_access_report(client, company, owner, make_user):
    make_user("cashier@scigate.test", "Cashier", )
    client.force_login(owner)
    page = client.get(reverse("users:access_report"))
    html = page.content.decode()
    assert page.status_code == 200 and "Cashier" in html and "Sara Owner" in html and "Review and sign-off" in html
    xlsx = client.get(reverse("users:access_report"), {"format": "xlsx"})
    assert xlsx["Content-Type"].startswith("application/vnd.openxmlformats")
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames[:2] == ["People", "Permissions"]


# ---------- 6. Delete drafts ----------
def test_delete_needs_permission(client, company, owner, make_user, g):
    rep = make_user("rep@scigate.test", "Sales officer")       # no sales.delete
    cust = Customer.objects.create(company=company, code="C-1", name_en="Nile", currency_id="EGP", account_manager=rep)
    draft = Invoice.objects.create(company=company, customer=cust, date=g["today"], due_date=g["today"],
                                   currency_id="EGP", created_by=rep)
    client.force_login(rep)
    assert client.post(reverse("core:doc_delete", args=["sales", "invoice", draft.pk])).status_code == 403
    assert Invoice.objects.filter(pk=draft.pk).exists()
    acc = make_user("acc@scigate.test", "Accountant")           # has sales.delete
    client.force_login(acc)
    client.post(reverse("core:doc_delete", args=["sales", "invoice", draft.pk]))
    assert not Invoice.objects.filter(pk=draft.pk).exists()


def test_existing_roles_upgraded(company):
    from users.defaults import upgrade_roles_batch_g

    custom = Role.objects.create(company=company, name_en="Branch sales", permissions=[
        "sales.view", "sales.create", "inventory.view", "purchases.approve"])
    upgrade_roles_batch_g(company)
    custom.refresh_from_db()
    assert {"sales.delete", "inventory.costs"} <= set(custom.permissions)
    assert ApprovalLimit.objects.get(role=custom, doc_type="purchase_order").max_amount >= D("999999999999")
    sales_officer = Role.objects.get(company=company, name_en="Sales officer")
    assert "inventory.costs" not in sales_officer.permissions and "sales.delete" not in sales_officer.permissions
