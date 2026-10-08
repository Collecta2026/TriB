"""Team and payroll screens: import staff from the payroll workbook, run payroll through its sign-offs, extract
deductions. All people here are invented."""
import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from banking.services import create_bank_account
from ledger.services import account_balance
from payroll.models import PayrollRun
from team.models import Employee

from .conftest import D, acct

PAGES = ["team:employees", "team:employee_new", "team:import", "team:departments", "team:benefits", "team:loans",
         "team:attendance", "team:leave", "payroll:runs", "payroll:settings", "payroll:leave_provision"]


def _workbook():
    import openpyxl
    from openpyxl.styles import PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Payroll September 2026"])
    ws.append(["Department", "SN", "Name", "Title", "Full Time", "Bank / Instapay", "Joining Date", "Net Basic Salary",
               "Allowance", "Monthly Employee Medical Due", "Monthly Employee Social Due"])
    ws.append(["Sales", 101, "Test Person A", "Sales rep", "F", "Bank", "2024-01-15", 12000, 1000, 250, 1100])
    ws.append(["Sales", 102, "Test Person B", "Driver", "P", "Instapay", "2025-03-01", 6000, 0, 0, 0])
    ws.append(["Warehouse", 103, "Test Person C", "Storekeeper", "F", "Cash", "2023-06-01", 8000, 500, 0, 880])
    for cell in ws[5]:
        cell.fill = PatternFill("solid", fgColor="FFFFFF00")  # yellow = on hold
    ws2 = wb.create_sheet("Other")
    ws2.append(["Name", "SN"])
    ws2.append(["Ignored", 999])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("lang", ["en", "ar"])
def test_team_and_payroll_pages_render(client, company, owner, lang):
    client.force_login(owner)
    client.cookies["django_language"] = lang
    for name in PAGES:
        assert client.get(reverse(name)).status_code == 200, name


def test_import_preview_then_confirm(client, company, owner):
    client.force_login(owner)
    upload = SimpleUploadedFile("Payroll Sep 2026.xlsx", _workbook(),
                                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    page = client.post(reverse("team:import"), {"workbook": upload}).content.decode()
    assert "Test Person A" in page and not Employee.objects.exists()
    client.post(reverse("team:import"), {"confirm": "1"})
    assert Employee.objects.count() == 3  # the second sheet is ignored
    a = Employee.objects.get(code="101")
    assert (a.department.name_en, a.basic_salary, a.allowances, a.medical_employee_share) == ("Sales", 12000, 1000, 250)
    assert a.insured and a.insurable_wage == D(10000)
    b = Employee.objects.get(code="102")
    assert b.contract_type == "part_time" and b.payment_method == "instapay" and not b.insured
    assert Employee.objects.get(code="103").status == "hold"


def test_payroll_run_signoffs_posting_and_extract(client, company, owner, make_user, cash):
    client.force_login(owner)
    client.post(reverse("team:import"), {"workbook": SimpleUploadedFile("p.xlsx", _workbook())})
    client.post(reverse("team:import"), {"confirm": "1"})
    response = client.post(reverse("payroll:runs"), {"year": "2026", "month": "9"})
    run = PayrollRun.objects.get()
    assert response.status_code == 302 and run.payslips.count() == 2  # the person on hold is not paid
    client.post(reverse("payroll:run_action", args=[run.pk]), {"action": "submit"})
    fm = make_user("fm@scigate.test", "Finance manager")
    md = make_user("md@scigate.test", "Managing director")
    client.force_login(fm)
    client.post(reverse("payroll:run_action", args=[run.pk]), {"action": "approve"})
    run.refresh_from_db()
    assert run.status == "approved" and run.journal_entry is None
    client.force_login(md)
    client.post(reverse("payroll:run_action", args=[run.pk]), {"action": "authorise"})
    run.refresh_from_db()
    assert run.status == "authorised" and run.journal_entry is not None
    t = run.totals()
    assert t["net_pay"] > 0
    assert account_balance(acct(company, "2155")) == t["net_pay"]  # salaries payable
    client.force_login(owner)
    xlsx = client.get(reverse("payroll:deductions", args=[run.pk]) + "?format=xlsx")
    assert xlsx.status_code == 200 and xlsx["Content-Type"].startswith("application/vnd.openxml")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Deductions", "NOSI", "Salary tax", "Payments", "GL journal"]
    gl = list(wb["GL journal"].iter_rows(min_row=2, values_only=True))
    assert round(sum(r[3] or 0 for r in gl), 2) == round(sum(r[4] or 0 for r in gl), 2)
    bank = create_bank_account(company, kind="bank", currency="EGP", name_en="CIB", user=owner)
    client.post(reverse("payroll:run_action", args=[run.pk]),
                {"action": "pay", "pay_date": "2026-09-30", "bank_account": bank.pk})
    run.refresh_from_db()
    assert run.status == "paid" and account_balance(acct(company, "2155")) == 0
    client.post(reverse("payroll:run_action", args=[run.pk]),
                {"action": "remit", "kind": "si", "pay_date": "2026-10-15", "bank_account": bank.pk, "reference": "NOSI"})
    assert run.remittances.get().amount == t["si_employee"] + t["si_employer"] + t["emergency_fund"]
    assert account_balance(acct(company, "2160")) == 0
    assert client.get(reverse("payroll:payslip", args=[run.payslips.first().pk])).status_code == 200
    assert client.get(reverse("payroll:payslips_print", args=[run.pk])).status_code == 200
    assert client.get(reverse("team:employee", args=[Employee.objects.get(code="101").pk])).status_code == 200
