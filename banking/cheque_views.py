"""Returned (bounced) cheques reported by the bank, and the cheque reconciliation report."""
import io
from decimal import Decimal

from django import forms
from django.contrib import messages
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy
from django.views.decorators.http import require_POST

from core.forms_util import DATE
from core.models import AuditLog, client_ip
from ledger.services import PostingError, account_balance, foreign_balance
from users.permissions import require_perm

from . import feeds, services
from .models import BankAccount, Cheque, ChequeEvent, StatementLine

SESSION_KEY = "cheques.bounce"
ZERO = Decimal("0")


class BounceNoticeForm(forms.Form):
    date = forms.DateField(label=gettext_lazy("Date of the bank's notice"), widget=DATE)
    bank_account = forms.ModelChoiceField(BankAccount.objects.none(), label=gettext_lazy("Bank"), required=False,
                                          help_text=gettext_lazy("Leave empty to look in every bank account."))
    reference = forms.CharField(label=gettext_lazy("Bank's reference"), max_length=80, required=False)
    numbers = forms.CharField(label=gettext_lazy("Returned cheques"), required=False,
                              widget=forms.Textarea(attrs={"rows": 6, "placeholder": "10048832, insufficient funds\n55120917"}),
                              help_text=gettext_lazy("One per line: cheque number, then a comma and the reason if the bank gave one."))
    file = forms.FileField(label=gettext_lazy("Or the bank's list (.csv or .xlsx)"), required=False,
                           help_text=gettext_lazy("First column: cheque number. Second column (optional): reason."))

    def __init__(self, *args, company, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["bank_account"].queryset = BankAccount.objects.filter(company=company, kind="bank", is_active=True)

    def clean(self):
        data = super().clean()
        if not data.get("numbers") and not data.get("file"):
            raise forms.ValidationError(_("Type the cheque numbers or upload the bank's list."))
        return data


def _parse_notice(data):
    rows = []
    if data.get("file"):
        try:
            table = feeds.read_table(data["file"])
        except Exception:  # damaged files raise many different errors
            raise forms.ValidationError(_("TriB could not read this file. Export it again as CSV or .xlsx."))
        rows += [(str(r[0]).strip(), str(r[1]).strip() if len(r) > 1 else "") for r in table if r and str(r[0]).strip()]
    for raw in (data.get("numbers") or "").splitlines():
        if raw.strip():
            number, _sep, reason = raw.partition(",")
            rows.append((number.strip(), reason.strip()))
    out, seen = [], set()
    for number, reason in rows:
        digits = number.lstrip("#").strip()
        if digits and any(ch.isdigit() for ch in digits) and digits not in seen:
            seen.add(digits)
            out.append((digits, reason[:200]))
    return out


def _find(company, number, bank_account):
    qs = Cheque.objects.filter(company=company, direction="in")
    found = [c for c in qs.filter(number__endswith=number.lstrip("0") or number) if c.number.lstrip("0") == number.lstrip("0")]
    if bank_account:
        found = [c for c in found if c.bank_account_id in (bank_account.pk, None)]
    return found


@require_perm("cheques.view")
def bounced(request):
    company = request.company
    form = BounceNoticeForm(request.POST or None, request.FILES or None, company=company,
                            initial={"date": timezone.localdate()})
    preview = None
    if request.method == "POST" and "confirm" in request.POST and request.membership.has_perm("cheques.post"):
        notice = request.session.pop(SESSION_KEY, None)
        if not notice:
            messages.error(request, _("The preview has expired. Enter the list again."))
            return redirect("banking:bounced")
        chosen = set(request.POST.getlist("cheque"))
        done = 0
        date = parse_date(notice["date"])
        for item in notice["rows"]:
            if str(item["cheque"]) not in chosen:
                continue
            cheque = get_object_or_404(Cheque, pk=item["cheque"], company=company)
            try:
                with transaction.atomic():
                    services.bounce_cheque(cheque, date, request.user, reason=item["reason"],
                                           reference=notice["reference"])
                done += 1
            except PostingError as exc:
                messages.error(request, f"#{cheque.number}: {exc}")
        if done:
            AuditLog.record(company, request.user, "cheque.bounce_notice", None,
                            f"{done} cheque(s) {notice['reference']}", ip=client_ip(request))
            messages.success(request, _("%(n)s cheques recorded as returned. The amounts are owed again by the customers.") % {"n": done})
        return redirect("banking:bounced")
    if request.method == "POST" and "preview" in request.POST:
        if not request.membership.has_perm("cheques.post"):
            messages.error(request, _("You do not have permission to do that."))
        elif form.is_valid():
            try:
                rows = _parse_notice(form.cleaned_data)
            except forms.ValidationError as exc:
                form.add_error("file", exc)
                rows = []
            d = form.cleaned_data
            preview = []
            for number, reason in rows:
                found = _find(company, number, d["bank_account"])
                open_ = [c for c in found if c.status in ("under_collection", "in_safe")]
                if len(open_) == 1:
                    preview.append({"number": number, "reason": reason, "cheque": open_[0], "state": "ok"})
                elif len(open_) > 1:
                    preview.append({"number": number, "reason": reason, "matches": open_, "state": "several"})
                elif found:
                    preview.append({"number": number, "reason": reason, "cheque": found[-1], "state": "closed"})
                else:
                    preview.append({"number": number, "reason": reason, "state": "missing"})
            request.session[SESSION_KEY] = {
                "date": d["date"].isoformat(), "reference": d["reference"],
                "rows": [{"cheque": r["cheque"].pk, "reason": r["reason"]} for r in preview if r["state"] == "ok"]}
    open_returned = (Cheque.objects.filter(company=company, direction="in", status="bounced")
                     .select_related("customer", "currency", "bank_account").order_by("bounced_on"))
    history = (ChequeEvent.objects.filter(company=company, action__in=("bounced", "resubmitted", "settled"))
               .select_related("cheque", "bank_account", "journal_entry").order_by("-date", "-id")[:30])
    today = timezone.localdate()
    accounts = BankAccount.objects.filter(company=company, is_active=True)
    return render(request, "banking/bounced.html", {
        "form": form, "preview": preview, "returned": open_returned, "history": history, "today": today,
        "banks": accounts.filter(kind="bank"), "money_accounts": accounts,
        "total": sum((c.amount for c in open_returned), ZERO)})


@require_POST
@require_perm("cheques.post")
def returned_action(request, pk):
    cheque = get_object_or_404(Cheque, pk=pk, company=request.company)
    account = BankAccount.objects.filter(company=request.company, pk=request.POST.get("bank_account") or 0).first()
    date = parse_date(request.POST.get("date") or "") or timezone.localdate()
    reference = (request.POST.get("reference") or "").strip()
    try:
        if request.POST.get("action") == "resubmit":
            services.resubmit_cheque(cheque, account, date, request.user, reference)
            messages.success(request, _("Cheque %(n)s is back with the bank for collection.") % {"n": cheque.number})
        elif request.POST.get("action") == "settle":
            services.settle_returned_cheque(cheque, account, date, request.user, reference)
            messages.success(request, _("Returned cheque %(n)s settled and the customer's account credited.") % {"n": cheque.number})
    except PostingError as exc:
        messages.error(request, str(exc))
    return redirect("banking:bounced")


# ---------- Cheque reconciliation report ----------
def _report(company, account, date_from, date_to, statement_balance):
    cheques = Cheque.objects.filter(company=company).select_related("currency", "customer")
    on_account = cheques.filter(bank_account=account)
    cleared = (ChequeEvent.objects.filter(company=company, action="cleared", bank_account=account,
                                          date__gte=date_from, date__lte=date_to)
               .select_related("cheque", "statement_line", "journal_entry").order_by("date", "id"))
    cleared_out = [e for e in cleared if e.cheque.direction == "out"]
    cleared_in = [e for e in cleared if e.cheque.direction == "in"]
    outstanding = list(on_account.filter(direction="out", status="issued", issue_date__lte=date_to).order_by("due_date"))
    collection = list(on_account.filter(direction="in", status="under_collection").order_by("due_date"))
    in_safe = list(cheques.filter(direction="in", status="in_safe", issue_date__lte=date_to).order_by("due_date"))
    returned = list(ChequeEvent.objects.filter(company=company, action="bounced", date__gte=date_from, date__lte=date_to)
                    .filter(Q(bank_account=account) | Q(bank_account__isnull=True)).select_related("cheque").order_by("date"))
    foreign = account.currency_id != company.base_currency_id
    book = foreign_balance(account.gl_account, account.currency_id, date_to) if foreign else \
        account_balance(account.gl_account, date_to)
    unbooked = list(StatementLine.objects.filter(bank_account=account, status="new", date__lte=date_to))
    unbooked_total = sum((l.amount for l in unbooked), ZERO)
    total = lambda items: sum((c.amount for c in items), ZERO)  # noqa: E731
    summary = {
        "book": book, "statement": statement_balance, "unbooked": unbooked_total,
        "difference": (statement_balance - unbooked_total - book) if statement_balance is not None else None,
        "outstanding": total(outstanding), "collection": total(collection), "in_safe": total(in_safe),
        "projected": book - total(outstanding) + total(collection),
        "cleared_out": sum((e.amount for e in cleared_out), ZERO), "cleared_in": sum((e.amount for e in cleared_in), ZERO),
        "returned": sum((e.amount for e in returned), ZERO),
    }
    return {"cleared_out": cleared_out, "cleared_in": cleared_in, "outstanding": outstanding, "collection": collection,
            "in_safe": in_safe, "returned": returned, "unbooked": unbooked, "s": summary}


@require_perm("cheques.view")
def cheque_reconciliation(request):
    company = request.company
    accounts = list(BankAccount.objects.filter(company=company, kind="bank", is_active=True))
    aid = request.GET.get("account", "")
    account = next((a for a in accounts if str(a.pk) == aid), accounts[0] if accounts else None)
    today = timezone.localdate()
    date_to = parse_date(request.GET.get("to", "") or "") or today
    date_from = parse_date(request.GET.get("from", "") or "") or date_to.replace(day=1)
    raw = (request.GET.get("statement_balance") or "").replace(",", "").strip()
    try:
        statement_balance = Decimal(raw) if raw else None
    except ArithmeticError:
        statement_balance = None
    ctx = {"accounts": accounts, "account": account, "from": date_from, "to": date_to, "raw_balance": raw, "today": today}
    if account:
        ctx.update(_report(company, account, date_from, date_to, statement_balance))
        if request.GET.get("format") == "xlsx" and request.membership.has_perm("cheques.view"):
            return _excel(account, date_from, date_to, ctx)
    return render(request, "banking/cheque_reconciliation.html", ctx)


def _excel(account, date_from, date_to, r):
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    def sheet(ws, title, header, rows):
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = bold
        for row in rows:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
        for i in range(1, len(header) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 22

    s = r["s"]
    sheet(wb.active, "Summary", ["Item", "Amount"], [
        ["Bank account", account.name_en], ["Period", f"{date_from} to {date_to}"],
        ["Balance per bank statement", s["statement"]], ["Statement lines not yet in the books", s["unbooked"]],
        ["Balance per books", s["book"]], ["Difference", s["difference"]],
        ["Issued cheques cleared", s["cleared_out"]], ["Received cheques cleared", s["cleared_in"]],
        ["Outstanding issued cheques", s["outstanding"]], ["Received cheques under collection", s["collection"]],
        ["Received cheques in the safe (not deposited)", s["in_safe"]], ["Cheques returned unpaid", s["returned"]],
        ["Projected balance after outstanding cheques", s["projected"]]])
    sheet(wb.create_sheet(), "Cleared", ["Cheque", "Direction", "Party", "Cheque date", "Cleared on", "Statement ref", "Amount"],
          [[e.cheque.number, e.cheque.get_direction_display(), e.cheque.party_name, e.cheque.issue_date, e.date,
            e.statement_line.reference if e.statement_line else "", e.amount] for e in r["cleared_out"] + r["cleared_in"]])
    sheet(wb.create_sheet(), "Outstanding issued", ["Cheque", "Payee", "Cheque date", "Due date", "Days out", "Amount"],
          [[c.number, c.party_name, c.issue_date, c.due_date, (date_to - c.issue_date).days, c.amount] for c in r["outstanding"]])
    sheet(wb.create_sheet(), "Under collection", ["Cheque", "Customer", "Drawn on", "Due date", "Amount"],
          [[c.number, c.party_name, c.drawee_bank, c.due_date, c.amount] for c in r["collection"]])
    sheet(wb.create_sheet(), "In safe", ["Cheque", "Customer", "Drawn on", "Due date", "Amount"],
          [[c.number, c.party_name, c.drawee_bank, c.due_date, c.amount] for c in r["in_safe"]])
    sheet(wb.create_sheet(), "Returned", ["Cheque", "Customer", "Returned on", "Reason", "Bank ref", "Amount"],
          [[e.cheque.number, e.cheque.party_name, e.date, e.note, e.reference, e.amount] for e in r["returned"]])
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="cheque-reconciliation-{date_to}.xlsx"'
    return response
