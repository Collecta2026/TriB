"""Import products and services from Excel (or CSV): one row per product. A row whose SKU already exists updates
that product; other rows add new ones (an empty SKU is numbered from the category). Nothing is saved until the
preview is confirmed. Categories that don't exist yet are created."""
import csv
import io
import re
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils.translation import gettext as _

COLUMNS = ["SKU", "Name", "Name (Arabic)", "Type", "Category", "Unit", "Sales description", "Purchase description",
           "Sales price", "Purchase cost", "Reorder level", "Barcode"]
MAX_ROWS = 3000
TYPE_WORDS = {"inventory": "inventory", "stock": "inventory", "stock item": "inventory", "مخزني": "inventory",
              "non_inventory": "non_inventory", "non-stock": "non_inventory", "non stock": "non_inventory",
              "non-stock item": "non_inventory", "غير مخزني": "non_inventory",
              "service": "service", "services": "service", "خدمة": "service"}


def template_workbook():
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(["", "Pallet wrap 500 mm", "", "Stock", "Packaging", "roll", "Stretch film, 23 micron", "", 12.5, 7.8, 20, ""])
    ws.append(["", "Customs clearance", "", "Service", "Services", "job", "Import customs clearance", "", 95, "", "", ""])
    for col, width in zip("ABCDEFGHIJKL", (12, 30, 24, 12, 16, 8, 32, 32, 12, 12, 12, 16)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rows(upload):
    if upload.name.lower().endswith(".csv"):
        text = upload.read().decode("utf-8-sig", errors="replace")
        return list(csv.reader(io.StringIO(text)))
    from openpyxl import load_workbook

    wb = load_workbook(upload, read_only=True, data_only=True)
    return [list(r) for r in wb.active.iter_rows(values_only=True)]


def _dec(v):
    if v in (None, ""):
        return None
    try:
        d = Decimal(str(v).replace(",", "").strip())
    except InvalidOperation:
        return "bad"
    return "bad" if d < 0 else d


def parse(company, upload):
    """(rows, problems): rows are plain dicts kept in the session until confirmed."""
    from .models import Item

    raw = _rows(upload)
    if not raw:
        return [], [_("The file is empty.")]
    head = [str(h or "").strip().lower() for h in raw[0]]
    col = {name: head.index(name.lower()) for name in COLUMNS if name.lower() in head}
    if "Name" not in col:
        return [], [_("The first row must hold the column names, including Name.")]
    existing = {i.sku: i.pk for i in Item.objects.filter(company=company)}
    rows, problems, seen = [], [], set()
    for n, r in enumerate(raw[1:MAX_ROWS + 1], start=2):
        def cell(name):
            i = col.get(name)
            v = r[i] if i is not None and i < len(r) else None
            return "" if v is None else str(v).strip() if not isinstance(v, (int, float)) else v
        name = str(cell("Name") or "").strip()
        if not name and not any(str(c or "").strip() for c in r):
            continue
        if not name:
            problems.append(_("Row %(n)s: the name is missing.") % {"n": n})
            continue
        sku = str(cell("SKU") or "").strip().upper()
        if sku and sku in seen:
            problems.append(_("Row %(n)s: SKU %(sku)s appears twice in the file.") % {"n": n, "sku": sku})
            continue
        seen.add(sku)
        kind = TYPE_WORDS.get(str(cell("Type") or "stock").strip().lower())
        if kind is None:
            problems.append(_("Row %(n)s: the type must be Stock, Non-stock or Service.") % {"n": n})
            continue
        numbers = {}
        for label, key in (("Sales price", "price"), ("Purchase cost", "cost"), ("Reorder level", "reorder")):
            value = _dec(cell(label))
            if value == "bad":
                problems.append(_("Row %(n)s: %(col)s is not a number.") % {"n": n, "col": label})
                break
            numbers[key] = None if value is None else str(value)
        else:
            rows.append({"row": n, "sku": sku, "update": existing.get(sku) if sku else None, "name": name[:200],
                         "name_ar": str(cell("Name (Arabic)") or "")[:200], "type": kind,
                         "category": str(cell("Category") or "")[:100], "unit": str(cell("Unit") or "")[:20],
                         "desc": str(cell("Sales description") or ""), "pdesc": str(cell("Purchase description") or ""),
                         "barcode": str(cell("Barcode") or "")[:60], **numbers})
    return rows, problems


def _category(company, name, cache):
    from .models import Category

    if not name:
        return None
    key = name.lower()
    if key not in cache:
        cat = Category.objects.filter(company=company, name_en__iexact=name).first() \
            or Category.objects.filter(company=company, prefix__iexact=name).first()
        if cat is None:
            base = (re.sub(r"[^A-Za-z]", "", name).upper() or "CAT")[:3]
            prefix, i = base, 1
            while Category.objects.filter(company=company, prefix=prefix).exists():
                i += 1
                prefix = f"{base[:4]}{i}"[:6]
            cat = Category.objects.create(company=company, name_en=name, name_ar=name, prefix=prefix)
        cache[key] = cat
    return cache[key]


@transaction.atomic
def save(company, rows):
    """Create or update the products. Returns (created, updated)."""
    from ledger.models import TaxRate

    from . import services
    from .models import Item

    tax = TaxRate.objects.filter(company=company, is_default=True, is_active=True).first()
    cache, created, updated = {}, 0, 0
    for r in rows:
        item = Item.objects.filter(company=company, pk=r["update"]).first() if r.get("update") else None
        new = item is None
        if new:
            item = Item(company=company, sku=r["sku"], sales_tax=tax, purchase_tax=tax)
        item.name_en, item.type = r["name"], r["type"]
        if r["name_ar"] or new:
            item.name_ar = r["name_ar"] or r["name"]
        if r["category"]:
            item.category = _category(company, r["category"], cache)
        for key, attr in (("unit", "unit"), ("desc", "description"), ("pdesc", "purchase_description"),
                          ("barcode", "barcode")):
            if r[key]:
                setattr(item, attr, r[key])
        for key, attr in (("price", "sales_price"), ("cost", "purchase_cost"), ("reorder", "reorder_level")):
            if r.get(key) is not None:
                setattr(item, attr, Decimal(r[key]))
        if item.type != "inventory":
            item.tracking = "none"
        if not item.sku:
            item.sku = services.next_sku(company, item.category)
        item.save()
        created, updated = created + new, updated + (not new)
    return created, updated
