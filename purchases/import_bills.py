"""Import bills from Excel (or CSV): one row per bill line; rows with the same supplier and supplier invoice number
become one bill. Nothing is saved until the preview is confirmed, and bills are created as drafts to review and post."""
import csv
import io
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils.translation import gettext as _

COLUMNS = ["Supplier", "Supplier invoice no", "Date", "Due date", "Product SKU", "Account code", "Description", "Qty",
           "Rate", "VAT %"]
MAX_ROWS = 2000


def template_workbook():
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(["SUP-00001", "INV-5521", date.today().strftime("%d/%m/%Y"), "", "", "6110", "Office rent October", 1, 25000, 14])
    ws.append(["SUP-00001", "INV-5521", date.today().strftime("%d/%m/%Y"), "", "", "6120", "Service charge", 1, 1500, 14])
    for col, width in zip("ABCDEFGHIJ", (16, 20, 12, 12, 14, 14, 34, 8, 12, 8)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rows(upload):
    name = upload.name.lower()
    if name.endswith(".csv"):
        text = upload.read().decode("utf-8-sig", errors="replace")
        return list(csv.reader(io.StringIO(text)))
    from openpyxl import load_workbook

    wb = load_workbook(upload, read_only=True, data_only=True)
    return [list(r) for r in wb.active.iter_rows(values_only=True)]


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _dec(v, default=None):
    if v in (None, ""):
        return default
    try:
        return Decimal(str(v).replace(",", "").strip())
    except InvalidOperation:
        return None


def parse(company, upload):
    """Returns (bills, problems). bills: [{supplier_id, supplier, invoice_no, date, due, lines: [...]}]."""
    from contacts.models import Supplier
    from inventory.models import Item
    from ledger.models import Account, TaxRate

    rows = _rows(upload)
    if not rows:
        return [], [_("The file is empty.")]
    header = [str(h or "").strip().lower() for h in rows[0]]
    want = [c.lower() for c in COLUMNS]
    missing = [COLUMNS[i] for i, c in enumerate(want) if c not in header]
    if missing:
        return [], [_("Missing columns: %(c)s. Download the template to see the layout.") % {"c": ", ".join(missing)}]
    idx = {c: header.index(c) for c in want}
    suppliers = {}
    for s in Supplier.objects.filter(company=company, is_active=True):
        suppliers[s.code.lower()] = s
        suppliers[s.name_en.lower()] = s
        if s.name_ar:
            suppliers[s.name_ar.lower()] = s
    items = {i.sku.lower(): i for i in Item.objects.filter(company=company, is_active=True)}
    accounts = {a.code: a for a in Account.objects.filter(company=company, is_group=False, is_active=True)}
    taxes = {t.rate: t for t in TaxRate.objects.filter(company=company, is_active=True)}
    bills, problems = {}, []
    for n, r in enumerate(rows[1:MAX_ROWS + 1], start=2):
        get = lambda col: r[idx[col.lower()]] if idx[col.lower()] < len(r) else None  # noqa: E731
        if not any(v not in (None, "") for v in r):
            continue
        errs = []
        sup = suppliers.get(str(get("Supplier") or "").strip().lower())
        if sup is None:
            errs.append(_("unknown supplier %(s)s") % {"s": get("Supplier")})
        inv_no = str(get("Supplier invoice no") or "").strip()
        day = _date(get("Date"))
        if day is None:
            errs.append(_("date missing or not dd/mm/yyyy"))
        due = _date(get("Due date"))
        sku = str(get("Product SKU") or "").strip()
        item = items.get(sku.lower()) if sku else None
        if sku and item is None:
            errs.append(_("unknown product %(s)s") % {"s": sku})
        if item is not None and item.is_stocked:
            errs.append(_("%(s)s is a stock item; receive it with an item receipt instead") % {"s": sku})
        code = str(get("Account code") or "").strip()
        if code.endswith(".0"):
            code = code[:-2]
        account = accounts.get(code) if code else None
        if code and account is None:
            errs.append(_("unknown account %(c)s") % {"c": code})
        if not item and not account:
            errs.append(_("give a product SKU or an account code"))
        qty, rate_ = _dec(get("Qty"), Decimal("1")), _dec(get("Rate"))
        if qty is None or qty <= 0:
            errs.append(_("quantity must be more than zero"))
        if rate_ is None:
            errs.append(_("rate missing"))
        vat = _dec(get("VAT %"), Decimal("0"))
        tax = taxes.get(vat) if vat else None
        if vat and tax is None:
            errs.append(_("no VAT rate of %(v)s%% is set up") % {"v": vat})
        if errs:
            problems.append(_("Row %(n)s: %(e)s") % {"n": n, "e": "; ".join(str(e) for e in errs)})
            continue
        key = (sup.pk, inv_no or f"row{n}")
        bill = bills.setdefault(key, {"supplier_id": sup.pk, "supplier": sup.name, "invoice_no": inv_no,
                                      "date": day.isoformat(), "due": (due or day + timedelta(days=sup.payment_terms_days)).isoformat(),
                                      "currency": sup.currency_id, "lines": [], "total": Decimal("0")})
        net = qty * rate_
        bill["lines"].append({"item": item.pk if item else None, "account": account.pk if account else None,
                              "description": str(get("Description") or "")[:300], "qty": str(qty), "rate": str(rate_),
                              "tax": tax.pk if tax else None})
        bill["total"] += net + (net * tax.rate / 100 if tax else 0)
    out = list(bills.values())
    for b in out:
        b["total"] = str(b["total"].quantize(Decimal("0.01")))
    return out, problems


@transaction.atomic
def create(company, bills, user):
    from purchases.models import Bill, BillLine

    made = []
    for b in bills:
        bill = Bill.objects.create(company=company, supplier_id=b["supplier_id"], supplier_invoice_no=b["invoice_no"],
                                   date=date.fromisoformat(b["date"]), due_date=date.fromisoformat(b["due"]),
                                   currency_id=b["currency"], created_by=user, reference=_("Imported"))
        for i, l in enumerate(b["lines"], start=1):
            BillLine.objects.create(document=bill, item_id=l["item"], account_id=l["account"], description=l["description"],
                                    qty=Decimal(l["qty"]), unit_price=Decimal(l["rate"]), tax_rate_id=l["tax"], position=i)
        made.append(bill)
    return made
