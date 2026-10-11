"""Turn the tables on a rendered report page into an Excel workbook, so every report exports the same way."""
import io
import re
from html.parser import HTMLParser

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

NUMBER = re.compile(r"^\(?-?[\d,]+(\.\d+)?\)?%?$")
SKIP_CLASSES = {"grid", "no-export"}


class _Tables(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables, self.title, self.sub = [], "", ""
        self._table = self._row = self._cell = None
        self._depth = 0
        self._in_h1 = self._in_sub = False
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        classes = set((a.get("class") or "").split())
        if tag == "h1" and not self.title:
            self._in_h1 = True
        elif tag == "div" and "sub" in classes and not self.sub:
            self._in_sub = True
        if tag == "table":
            self._depth += 1
            if self._depth == 1 and "tbl" in classes and not classes & SKIP_CLASSES:
                self._table = []
        elif self._table is not None and self._depth == 1:
            if tag == "tr":
                self._row = {"cells": [], "head": False, "cls": classes}
            elif tag in ("td", "th") and self._row is not None:
                self._cell = {"text": "", "span": int(a.get("colspan") or 1), "head": tag == "th"}
            elif tag in ("button", "form", "select", "script"):
                self._skip += 1

    def handle_endtag(self, tag):
        if tag == "h1":
            self._in_h1 = False
        elif tag == "div":
            self._in_sub = False
        if tag == "table":
            if self._depth == 1 and self._table is not None:
                self.tables.append(self._table)
                self._table = None
            self._depth -= 1
        elif self._table is not None and self._depth == 1:
            if tag in ("td", "th") and self._cell is not None:
                self._row["cells"].append(self._cell)
                self._row["head"] = self._row["head"] or self._cell["head"]
                self._cell = None
            elif tag == "tr" and self._row is not None:
                if self._row["cells"]:
                    self._table.append(self._row)
                self._row = None
            elif tag in ("button", "form", "select", "script") and self._skip:
                self._skip -= 1

    def handle_data(self, data):
        if self._in_h1:
            self.title += data
        elif self._in_sub:
            self.sub += data
        if self._cell is not None and not self._skip:
            self._cell["text"] += data


def _value(text):
    text = " ".join(text.split())
    if text and NUMBER.match(text):
        negative = text.startswith("(") or text.startswith("-")
        digits = text.strip("()%").replace(",", "").lstrip("-")
        try:
            number = float(digits)
        except ValueError:
            return text
        return -number if negative else number
    return text


def tables_to_workbook(html):
    parser = _Tables()
    parser.feed(html)
    title = " ".join(parser.title.split())
    wb = Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[\[\]:*?/\\]", "", title) or "Report")[:31]
    ws.sheet_view.rightToLeft = 'dir="rtl"' in html[:400]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    if parser.sub.strip():
        ws.append([" ".join(parser.sub.split())])
    widths = {}
    head_fill = PatternFill("solid", fgColor="E2EFF0")
    for table in parser.tables:
        ws.append([])
        for r in table:
            values = []
            for c in r["cells"]:
                values.append(_value(c["text"]))
                values.extend([""] * (c["span"] - 1))
            ws.append(values)
            bold = r["head"] or r["cls"] & {"grp", "sub", "tot"}
            for i, v in enumerate(values, start=1):
                cell = ws.cell(row=ws.max_row, column=i)
                if bold:
                    cell.font = Font(bold=True)
                if r["head"]:
                    cell.fill = head_fill
                if isinstance(v, float):
                    cell.number_format = "#,##0.00"
                widths[i] = max(widths.get(i, 8), min(len(str(v)) + 2, 60))
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), title
